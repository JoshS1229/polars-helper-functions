from datetime import date, datetime
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import polars as pl
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from polars_helper_functions import write_excel_polars
from polars_helper_functions._xlsx import _MAIN, _date_styles


@pytest.mark.parametrize("mode", ["raw", "table"])
@pytest.mark.parametrize("operation", ["new", "add", "replace"])
@pytest.mark.parametrize("unit", ["ms", "us", "ns"])
def test_date_formats_follow_column_types(tmp_path, mode, operation, unit):
    path = tmp_path / "dates.xlsx"
    if operation != "new":
        wb = Workbook()
        wb.active.title = "Keep"
        if operation == "replace":
            wb.create_sheet("Output")["A10"] = "stale"
        wb.save(path)
        wb.close()
    timestamp = datetime(2024, 2, 29, 15, 45, 30, 123000)
    frame = pl.DataFrame({
        "date": pl.Series([date(2024, 2, 29), None], dtype=pl.Date),
        "timestamp": pl.Series([timestamp, None], dtype=pl.Datetime(unit)),
        "null_date": pl.Series([None, None], dtype=pl.Date),
        "null_datetime": pl.Series([None, None], dtype=pl.Datetime(unit)),
        "text": ["2024-02-29", None],
        "number": [45351, None],
    })
    write_excel_polars(path, frame, "Output", mode=mode, table_name="Dates")
    wb = load_workbook(path)
    try:
        ws = wb["Output"]
        value = ws["A2"].value
        assert (value.date() if isinstance(value, datetime) else value) == date(2024, 2, 29)
        assert ws["B2"].value == timestamp
        for column, expected in [("A", "yyyy-mm-dd"), ("B", "yyyy-mm-dd hh:mm:ss"),
                                 ("C", "yyyy-mm-dd"), ("D", "yyyy-mm-dd hh:mm:ss")]:
            assert ws[f"{column}1"].number_format == "General"
            for row in [2, 3]:
                assert ws[f"{column}{row}"].number_format == expected
        assert ws["A2"].is_date and ws["B2"].is_date
        assert ws["A3"].value is None and ws["C2"].value is None
        assert ws["E2"].value == "2024-02-29"
        assert ws["E2"].number_format == ws["F2"].number_format == "General"
        assert ws["A10"].value is None
        if operation != "new":
            assert "Keep" in wb.sheetnames
        if mode == "table":
            assert "Dates" in ws.tables
    finally:
        wb.close()


def test_date_export_preserves_existing_styles_and_reuses_added_styles(tmp_path):
    path = tmp_path / "existing.xlsx"
    wb = Workbook()
    wb.epoch = CALENDAR_MAC_1904
    keep = wb.active
    keep.title = "Keep"
    keep["A1"] = 12.345
    keep["A1"].number_format = '0.0000 "custom"'
    keep["A1"].font = Font(bold=True, color="FF0000")
    keep["A1"].fill = PatternFill("solid", fgColor="00FF00")
    keep["B1"] = "=A1*2"
    wb.save(path)
    wb.close()
    with ZipFile(path) as archive:
        before = {name: archive.read(name) for name in archive.namelist()}
        original_styles = ET.fromstring(before["xl/styles.xml"])
    frame = pl.DataFrame({"date": [date(2025, 1, 2)], "dt": [datetime(2025, 1, 2, 3, 4, 5)]})
    count_after = None
    for name in ["Dates1", "Dates2", "Dates1"]:
        write_excel_polars(path, frame, name)
        with ZipFile(path) as archive:
            styles = ET.fromstring(archive.read("xl/styles.xml"))
            for tag in ["fonts", "fills", "borders", "cellStyleXfs", "cellStyles"]:
                assert ET.tostring(styles.find(f"{{{_MAIN}}}{tag}")) == ET.tostring(original_styles.find(f"{{{_MAIN}}}{tag}"))
            xfs = styles.find(f"{{{_MAIN}}}cellXfs")
            if count_after is None:
                count_after = len(xfs)
            assert len(xfs) == count_after
            for old, current in zip(original_styles.find(f"{{{_MAIN}}}cellXfs"), xfs):
                assert ET.tostring(old) == ET.tostring(current)
            assert archive.read("xl/worksheets/sheet1.xml") == before["xl/worksheets/sheet1.xml"]
    wb = load_workbook(path)
    try:
        assert wb.epoch == CALENDAR_MAC_1904
        assert wb["Keep"]["A1"].font.bold
        assert wb["Keep"]["A1"].fill.fgColor.rgb == "0000FF00"
        assert wb["Keep"]["A1"].number_format == '0.0000 "custom"'
        assert wb["Keep"]["B1"].value == "=A1*2"
        assert wb["Dates2"]["A2"].value == date(2025, 1, 2)
    finally:
        wb.close()


def test_style_extensions_and_prefixes_are_preserved():
    original = f'''<s:styleSheet xmlns:s="{_MAIN}"
        xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
        xmlns:x14="urn:example-extension" mc:Ignorable="x14">
        <s:cellXfs count="1"><s:xf numFmtId="0"/></s:cellXfs>
        <s:extLst><s:ext uri="custom"><x14:custom value="keep"/></s:ext></s:extLst>
    </s:styleSheet>'''.encode()
    updated, formats = _date_styles(original, ["yyyy-mm-dd"])
    assert b'xmlns:x14="urn:example-extension"' in updated
    assert b'mc:Ignorable="x14"' in updated
    root = ET.fromstring(updated)
    assert root.find(".//{urn:example-extension}custom").get("value") == "keep"
    assert formats["yyyy-mm-dd"] == 1


@pytest.mark.parametrize("existing", [False, True])
def test_timezone_columns_are_rejected_before_writing(tmp_path, existing):
    path = tmp_path / "dates.xlsx"
    if existing:
        Workbook().save(path)
    before = path.read_bytes() if existing else None
    frame = pl.DataFrame({"timestamp": pl.Series([None], dtype=pl.Datetime(time_zone="UTC"))})
    with pytest.raises(ValueError, match="Excel does not support time zones"):
        write_excel_polars(path, frame, "Output")
    assert (path.read_bytes() if path.exists() else None) == before


def test_adds_styles_to_workbook_without_styles_part(tmp_path):
    path = tmp_path / "no_styles.xlsx"
    Workbook().save(path)
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    del parts["xl/styles.xml"]
    rels = ET.fromstring(parts["xl/_rels/workbook.xml.rels"])
    for relation in list(rels):
        if relation.get("Type").endswith("/styles"):
            rels.remove(relation)
    parts["xl/_rels/workbook.xml.rels"] = ET.tostring(rels)
    types = ET.fromstring(parts["[Content_Types].xml"])
    for override in list(types):
        if override.get("PartName") == "/xl/styles.xml":
            types.remove(override)
    parts["[Content_Types].xml"] = ET.tostring(types)
    with ZipFile(path, "w") as archive:
        for name, contents in parts.items():
            archive.writestr(name, contents)

    write_excel_polars(path, pl.DataFrame({"date": [date(2025, 1, 2)]}), "Dates")
    wb = load_workbook(path)
    try:
        assert wb["Dates"]["A2"].value == date(2025, 1, 2)
        assert wb["Dates"]["A2"].number_format == "yyyy-mm-dd"
        assert wb["Dates"]["A2"].font is not None
    finally:
        wb.close()
