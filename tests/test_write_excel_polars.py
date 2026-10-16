from datetime import timedelta
from zipfile import ZipFile

import polars as pl
from openpyxl import Workbook, load_workbook

from polars_helper_functions import write_excel_polars


def test_replaces_only_target_worksheet_part(tmp_path):
    path = tmp_path / "existing.xlsx"
    workbook = Workbook()
    source = workbook.active
    source.title = "Source"
    source["A1"] = "old"
    formulas = workbook.create_sheet("Calculations")
    formulas["A1"] = "=Source!A2*2"
    workbook.save(path)

    with ZipFile(path, "a") as archive:
        archive.writestr("xl/pivotCache/pivotCacheDefinition99.xml", b"<opaque-pivot-data/>")

    with ZipFile(path) as archive:
        formula_before = archive.read("xl/worksheets/sheet2.xml")
        pivot_before = archive.read("xl/pivotCache/pivotCacheDefinition99.xml")

    write_excel_polars(path, pl.DataFrame({"value": [21]}), "Source")

    with ZipFile(path) as archive:
        assert archive.read("xl/worksheets/sheet2.xml") == formula_before
        assert archive.read("xl/pivotCache/pivotCacheDefinition99.xml") == pivot_before

    result = load_workbook(path, data_only=False)
    assert result["Source"]["A2"].value == 21
    assert result["Calculations"]["A1"].value == "=Source!A2*2"


def test_adds_sheet_and_can_write_table_to_existing_workbook(tmp_path):
    path = tmp_path / "existing.xlsx"
    Workbook().save(path)

    write_excel_polars(
        path,
        pl.DataFrame({"name": ["Alice", " Bob "]}),
        "Output",
        mode="table",
        table_name="OutputTable",
    )

    result = load_workbook(path)
    assert result["Output"]["A3"].value == " Bob "
    assert "OutputTable" in result["Output"].tables


def test_writes_timedelta_as_excel_day_fraction(tmp_path):
    path = tmp_path / "existing.xlsx"
    Workbook().save(path)

    write_excel_polars(
        path,
        pl.DataFrame({"elapsed": [timedelta(days=1, hours=12)]}),
        "Output",
    )

    result = load_workbook(path, data_only=True)
    assert result["Output"]["A2"].value == 1.5


def test_rejects_unknown_mode_before_creating_a_file(tmp_path):
    path = tmp_path / "new.xlsx"

    try:
        write_excel_polars(path, pl.DataFrame({"a": [1]}), "Output", mode="unknown")
    except ValueError as error:
        assert "mode" in str(error)
    else:
        raise AssertionError("Expected ValueError")

    assert not path.exists()


def test_writes_all_nine_duration_sheets(tmp_path):
    names = [
        "1 PE-Const Rare Mean", "2 PE-Complete Rare Mean",
        "3 PE-Const Unlinked Mean", "4 PE-Const All Linked Mean",
        "5 PE-Const Rare Median", "N PE-Const Rare", "N PE-Complete Rare",
        "N PE-Const Unlinked", "N PE-Const All Linked",
    ]
    path = tmp_path / "Project Duration Summary.xlsx"
    for index, name in enumerate(names):
        write_excel_polars(path, pl.DataFrame({"State": ["NY"], "2000": [index]}), name)

    result = load_workbook(path)
    try:
        assert result.sheetnames == names
        for index, name in enumerate(names):
            assert result[name]["B2"].value == index
    finally:
        result.close()


def test_retries_temporary_lock_after_closing_archives(tmp_path, monkeypatch):
    from polars_helper_functions import _xlsx

    path = tmp_path / "existing.xlsx"
    write_excel_polars(path, pl.DataFrame({"value": [1]}), "First")
    original_replace = _xlsx.os.replace
    archives = []
    sleeps = []
    attempts = []

    class TrackedZipFile(ZipFile):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            archives.append(self)

    def temporarily_locked(source, destination):
        assert all(archive.fp is None for archive in archives)
        attempts.append(source)
        if len(attempts) < 3:
            raise PermissionError(13, "Access is denied", str(destination))
        original_replace(source, destination)

    monkeypatch.setattr(_xlsx, "ZipFile", TrackedZipFile)
    monkeypatch.setattr(_xlsx.os, "replace", temporarily_locked)
    monkeypatch.setattr(_xlsx._time, "sleep", sleeps.append)

    write_excel_polars(path, pl.DataFrame({"value": [2]}), "N PE-Const Rare")

    assert len(attempts) == 3
    assert len(set(attempts)) == 1  # Retry the completed archive, not the serialization.
    assert sleeps == [0.1, 0.2]
    result = load_workbook(path)
    try:
        assert result["First"]["A2"].value == 1
        assert result["N PE-Const Rare"]["A2"].value == 2
    finally:
        result.close()
    assert list(tmp_path.iterdir()) == [path]


def test_persistent_lock_preserves_original_and_cleans_temp(tmp_path, monkeypatch):
    import pytest
    from polars_helper_functions import _xlsx

    path = tmp_path / "existing.xlsx"
    write_excel_polars(path, pl.DataFrame({"value": [1]}), "First")
    before = path.read_bytes()
    attempts = []
    sleeps = []

    def locked(source, destination):
        attempts.append(source)
        raise PermissionError(13, "Access is denied", str(destination))

    monkeypatch.setattr(_xlsx.os, "replace", locked)
    monkeypatch.setattr(_xlsx._time, "sleep", sleeps.append)

    with pytest.raises(PermissionError, match="Close the workbook in Excel") as raised:
        write_excel_polars(path, pl.DataFrame({"value": [2]}), "Second")

    assert isinstance(raised.value.__cause__, PermissionError)
    assert raised.value.filename == str(path)
    assert len(attempts) == 6
    assert sleeps == [0.1, 0.2, 0.4, 0.8, 1.6]
    assert path.read_bytes() == before
    assert list(tmp_path.iterdir()) == [path]


def test_other_replace_errors_are_not_retried(tmp_path, monkeypatch):
    import pytest
    from polars_helper_functions import _xlsx

    path = tmp_path / "existing.xlsx"
    Workbook().save(path)
    before = path.read_bytes()

    def broken(source, destination):
        raise OSError("disk failure")

    def unexpected_sleep(delay):
        raise AssertionError("Non-permission errors must not be retried")

    monkeypatch.setattr(_xlsx.os, "replace", broken)
    monkeypatch.setattr(_xlsx._time, "sleep", unexpected_sleep)
    with pytest.raises(OSError, match="disk failure"):
        write_excel_polars(path, pl.DataFrame({"value": [2]}), "Second")
    assert path.read_bytes() == before
    assert list(tmp_path.iterdir()) == [path]
